"""Create mis-styled DOCX test documents for heading-correction evaluation.

The no-style corpus (``create_no_style_docs.py``) covers documents with *no*
typographic signal.  This corpus covers the opposite failure: documents whose
typography is present but **wrong** — the author started at Heading 2, numbered
top-level sections ``1.1``, applied a heading style to a body paragraph, made
deep headings larger than shallow ones, or bolded the entire document.

Every variant of one logical structure shares the same ground truth: the *true*
logical hierarchy.  The rendered heading text may carry the author's wrong
numbering (that is the perturbation), but ``level`` is always the truth, so a
parser is scored on recovering the intended structure rather than echoing the
typography.

GT discipline (see ``tests/data/GT_README.md``):

- titles/levels come from this generator (author-side, independent of the system
  under test); only ``block_id`` is located by running the provider, exactly as
  ``evaluation/repair_gt.py`` locates TOC entries;
- ``block_id = -1`` marks "heading exists but could not be located" — never 0;
- ``distractors`` lists blocks that must **not** be reported as headings.

Run: python scripts/manual/create_mis_styled_docs.py
Outputs: tests/data/mis_styled/{structure}__{variant}.docx (+ .json)
"""
from __future__ import annotations

import json
import os
import sys
from dataclasses import dataclass, field

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..")))
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx import Document
from docx.shared import Pt

from infrastructure.providers.docx_provider import DocxProvider
from modules.parser.titles import normalize_title

OUT_DIR = os.path.join("tests", "data", "mis_styled")
BODY_SIZE = 11.0
# Font size by *true* level, used by the variants that convey depth via size.
LEVEL_SIZE = {1: 16.0, 2: 14.0, 3: 12.5, 4: 11.5}


@dataclass(frozen=True)
class Node:
    """One logical section: its true depth, title text, and body paragraphs."""

    level: int
    title: str
    body: tuple[str, ...]


@dataclass
class Para:
    """One rendered paragraph plus the GT role it plays."""

    text: str
    style: str | None = None
    size: float = BODY_SIZE
    bold: bool = False
    gt_level: int | None = None   # set => this paragraph is a true heading
    distractor: bool = False      # set => must never be reported as a heading


@dataclass
class Variant:
    """A typography perturbation applied to every structure."""

    name: str
    description: str
    render: object = field(repr=False)


# ── Logical structures (the ground truth, before any typography) ──────

_LOREM = (
    "This section records the operational context in enough detail that a "
    "reader can reconstruct the decision without consulting the appendix."
)
_LOREM2 = (
    "Measurements were repeated three times and the median is reported; the "
    "spread never exceeded four percent across the runs."
)
_CJK = "本节说明该阶段的输入、输出与失败模式，并给出可复现实验的参数取值。"
_CJK2 = "所有数值均为三次独立运行的中位数，波动范围不超过百分之四。"

STRUCTURES: dict[str, tuple[Node, ...]] = {
    "report": (
        Node(1, "Programme Overview", (_LOREM, _LOREM2)),
        Node(2, "Scope and Boundaries", (_LOREM,)),
        Node(2, "Known Risks", (_LOREM2, _LOREM)),
        Node(1, "Technical Design", (_LOREM,)),
        Node(2, "Storage Layer", (_LOREM2,)),
        Node(2, "Serving Layer", (_LOREM, _LOREM2)),
        Node(1, "Rollout Plan", (_LOREM2,)),
        Node(2, "Staged Release", (_LOREM,)),
        Node(2, "Rollback Criteria", (_LOREM2,)),
    ),
    "deep": (
        Node(1, "System Overview", (_LOREM,)),
        Node(2, "Architecture", (_LOREM2,)),
        Node(3, "Frontend Runtime", (_LOREM,)),
        Node(3, "Backend Runtime", (_LOREM2,)),
        Node(4, "Request Pipeline", (_LOREM,)),
        Node(4, "Persistence Adapters", (_LOREM2,)),
        Node(2, "Deployment Topology", (_LOREM,)),
        Node(1, "Data Pipeline", (_LOREM2,)),
        Node(2, "Ingestion", (_LOREM,)),
        Node(2, "Transformation", (_LOREM2,)),
        Node(3, "Validation Rules", (_LOREM,)),
        Node(3, "Normalisation", (_LOREM2,)),
        Node(1, "Observed Results", (_LOREM,)),
        Node(2, "Throughput", (_LOREM2,)),
        Node(2, "Accuracy", (_LOREM,)),
    ),
    "cjk_paper": (
        Node(1, "研究背景", (_CJK, _CJK2)),
        Node(2, "问题定义", (_CJK,)),
        Node(2, "相关工作", (_CJK2,)),
        Node(1, "方法设计", (_CJK,)),
        Node(2, "候选生成", (_CJK2,)),
        Node(2, "全局解码", (_CJK, _CJK2)),
        Node(1, "实验评估", (_CJK2,)),
        Node(2, "数据与指标", (_CJK,)),
        Node(2, "结果分析", (_CJK2,)),
        Node(1, "结论", (_CJK,)),
    ),
}


def _true_numbering(nodes: tuple[Node, ...]) -> list[str]:
    """Depth-correct section numbers: 1, 1.1, 1.1.1, 1.2, 2, ..."""
    counters = [0] * 6
    numbers: list[str] = []
    for node in nodes:
        counters[node.level - 1] += 1
        for deeper in range(node.level, 6):
            counters[deeper] = 0
        numbers.append(".".join(str(counters[i]) for i in range(node.level)))
    return numbers


# ── Variant renderers ─────────────────────────────────────────────────
#
# Each returns the full paragraph sequence of one rendered document.  The GT is
# derived from the ``gt_level`` / ``distractor`` flags, so a renderer cannot
# silently disagree with its own ground truth.

def _render_clean(nodes: tuple[Node, ...]) -> list[Para]:
    """Control: correct Heading N styles, depth-correct numbering."""
    paras: list[Para] = []
    for node, number in zip(nodes, _true_numbering(nodes), strict=True):
        paras.append(Para(
            text=f"{number} {node.title}",
            style=f"Heading {node.level}",
            size=LEVEL_SIZE[node.level],
            bold=True,
            gt_level=node.level,
        ))
        paras.extend(Para(text=body) for body in node.body)
    return paras


def _render_style_shift(nodes: tuple[Node, ...]) -> list[Para]:
    """Author started at Heading 2: every style is one level too deep."""
    paras: list[Para] = []
    for node, number in zip(nodes, _true_numbering(nodes), strict=True):
        shifted = min(node.level + 1, 6)
        paras.append(Para(
            text=f"{number} {node.title}",
            style=f"Heading {shifted}",
            size=LEVEL_SIZE[node.level],
            bold=True,
            gt_level=node.level,
        ))
        paras.extend(Para(text=body) for body in node.body)
    return paras


def _render_flat_style(nodes: tuple[Node, ...]) -> list[Para]:
    """One style for every depth: only numbering still carries the hierarchy."""
    paras: list[Para] = []
    for node, number in zip(nodes, _true_numbering(nodes), strict=True):
        paras.append(Para(
            text=f"{number} {node.title}",
            style="Heading 2",
            size=13.0,
            bold=True,
            gt_level=node.level,
        ))
        paras.extend(Para(text=body) for body in node.body)
    return paras


def _render_over_styled(nodes: tuple[Node, ...]) -> list[Para]:
    """Correct heading styles, but body paragraphs also carry a heading style."""
    paras: list[Para] = []
    body_ordinal = 0
    for node, number in zip(nodes, _true_numbering(nodes), strict=True):
        paras.append(Para(
            text=f"{number} {node.title}",
            style=f"Heading {node.level}",
            size=LEVEL_SIZE[node.level],
            bold=True,
            gt_level=node.level,
        ))
        for body in node.body:
            body_ordinal += 1
            # Every second body paragraph is mis-styled as a heading.  The text
            # is unchanged prose, so only shape — not wording — reveals it.
            if body_ordinal % 2 == 0:
                paras.append(Para(
                    text=body,
                    style="Heading 3",
                    size=BODY_SIZE,
                    distractor=True,
                ))
            else:
                paras.append(Para(text=body))
    return paras


def _render_numbering_shift(nodes: tuple[Node, ...]) -> list[Para]:
    """No heading styles; visible numbering is one level too deep (1.1 = top)."""
    paras: list[Para] = []
    for node, number in zip(nodes, _true_numbering(nodes), strict=True):
        paras.append(Para(
            text=f"1.{number} {node.title}",
            size=LEVEL_SIZE[node.level],
            bold=True,
            gt_level=node.level,
        ))
        paras.extend(Para(text=body) for body in node.body)
    return paras


def _render_font_inverted(nodes: tuple[Node, ...]) -> list[Para]:
    """No styles, no numbering, and deeper headings are rendered *larger*."""
    inverted = {1: 12.5, 2: 14.0, 3: 16.0, 4: 17.5}
    paras: list[Para] = []
    for node in nodes:
        paras.append(Para(
            text=node.title,
            size=inverted[node.level],
            bold=True,
            gt_level=node.level,
        ))
        paras.extend(Para(text=body) for body in node.body)
    return paras


_LABELS = (
    "Owner: Platform Team",
    "Status: In review",
    "Last updated: 2026-03-11",
    "Reviewers: A. Chen, B. Kumar",
)
_LABELS_CJK = (
    "负责人：平台组",
    "状态：评审中",
    "更新时间：2026-03-11",
    "评审人：陈某、库马尔",
)


def _labels_for(nodes: tuple[Node, ...]) -> tuple[str, ...]:
    """Pick label wording matching the structure's own language."""
    joined = "".join(node.title for node in nodes)
    if any("\u4e00" <= char <= "\u9fff" for char in joined):
        return _LABELS_CJK
    return _LABELS


def _render_styled_labels(nodes: tuple[Node, ...]) -> list[Para]:
    """Correct headings, plus short metadata lines wearing a heading style.

    The distractors here are deliberately **not** prose: they are short, they
    carry no sentence punctuation, and no amount of shape analysis can tell
    them from a real heading.  What gives them away is that they wear the
    *body's* exact typographic signature (body size, not bold) in a document
    whose real headings are all larger and bold — the author's own document
    draws the line, and these lines sit on the wrong side of it.
    """
    labels = _labels_for(nodes)
    paras: list[Para] = []
    for index, (node, number) in enumerate(
        zip(nodes, _true_numbering(nodes), strict=True),
    ):
        paras.append(Para(
            text=f"{number} {node.title}",
            style=f"Heading {node.level}",
            size=LEVEL_SIZE[node.level],
            bold=True,
            gt_level=node.level,
        ))
        paras.append(Para(
            text=labels[index % len(labels)],
            style="Heading 3",
            size=BODY_SIZE,
            distractor=True,
        ))
        paras.extend(Para(text=body) for body in node.body)
    return paras


def _render_bold_everything(nodes: tuple[Node, ...]) -> list[Para]:
    """Whole document bold at one size: the emphasis signal carries nothing."""
    paras: list[Para] = []
    for node, number in zip(nodes, _true_numbering(nodes), strict=True):
        paras.append(Para(
            text=f"{number} {node.title}",
            size=BODY_SIZE,
            bold=True,
            gt_level=node.level,
        ))
        paras.extend(
            Para(text=body, size=BODY_SIZE, bold=True) for body in node.body
        )
    return paras


VARIANTS: tuple[Variant, ...] = (
    Variant("clean", "Control: correct Heading N styles and depth-correct numbering", _render_clean),
    Variant("style_shift", "Every heading style one level too deep (document starts at Heading 2)", _render_style_shift),
    Variant("flat_style", "All depths share Heading 2; only numbering carries the hierarchy", _render_flat_style),
    Variant("over_styled", "Correct headings plus body paragraphs mis-styled as Heading 3", _render_over_styled),
    Variant("styled_labels", "Correct headings plus short metadata lines mis-styled as Heading 3 at body size", _render_styled_labels),
    Variant("numbering_shift", "No styles; visible numbering one level too deep (top sections are 1.1)", _render_numbering_shift),
    Variant("font_inverted", "No styles or numbering; deeper headings use a larger font", _render_font_inverted),
    Variant("bold_everything", "Entire document bold at body size; emphasis carries no information", _render_bold_everything),
)


# ── Rendering and GT location ─────────────────────────────────────────

def _write_docx(paras: list[Para], path: str) -> None:
    document = Document()
    for para in paras:
        paragraph = document.add_paragraph()
        if para.style:
            paragraph.style = document.styles[para.style]
        run = paragraph.add_run(para.text)
        run.font.size = Pt(para.size)
        run.font.bold = para.bold
    document.save(path)


def _normalise(text: str | None) -> str:
    """Compare on the structural name only.

    The DOCX provider renders bold runs as ``**…**`` for lossless transport, and
    every heading here is bold, so raw text comparison would never match.
    """
    return normalize_title(text)


def _locate_block_ids(paras: list[Para], path: str) -> tuple[list[int | None], int]:
    """Map each rendered paragraph to the block id the provider assigned it.

    Alignment is positional over text blocks and verified paragraph by
    paragraph; any paragraph whose text does not match its aligned block falls
    back to a unique full-document text search, and finally to ``None`` (which
    the caller writes as the ``-1`` sentinel).
    """
    blocks = [
        block for block in DocxProvider().extract(path)
        if block.type == "text" and block.text
    ]
    located: list[int | None] = []
    mismatches = 0
    for index, para in enumerate(paras):
        expected = _normalise(para.text)
        if index < len(blocks) and _normalise(blocks[index].text) == expected:
            located.append(blocks[index].id)
            continue
        hits = [block.id for block in blocks if _normalise(block.text) == expected]
        if len(hits) == 1:
            located.append(hits[0])
        else:
            located.append(None)
        mismatches += 1
    return located, mismatches


def main() -> None:
    os.makedirs(OUT_DIR, exist_ok=True)
    print(f"Creating mis-styled documents in {OUT_DIR} ...\n")

    manifest: list[dict] = []
    total_unlocated = 0

    for structure_name, nodes in STRUCTURES.items():
        for variant in VARIANTS:
            paras = variant.render(nodes)
            stem = f"{structure_name}__{variant.name}"
            docx_path = os.path.join(OUT_DIR, f"{stem}.docx")
            json_path = os.path.join(OUT_DIR, f"{stem}.json")

            _write_docx(paras, docx_path)
            block_ids, mismatches = _locate_block_ids(paras, docx_path)

            headings = [
                {
                    "block_id": block_id if block_id is not None else -1,
                    "title": para.text,
                    "level": para.gt_level,
                }
                for para, block_id in zip(paras, block_ids, strict=True)
                if para.gt_level is not None
            ]
            distractors = [
                {
                    "block_id": block_id if block_id is not None else -1,
                    "text": para.text[:80],
                }
                for para, block_id in zip(paras, block_ids, strict=True)
                if para.distractor
            ]
            unlocated = sum(1 for h in headings if h["block_id"] < 0)
            total_unlocated += unlocated

            gt = {
                "doc_title": "",
                "doc_authors": "",
                "gt_source": "synthetic_perturbation",
                "structure": structure_name,
                "variant": variant.name,
                "perturbation": variant.description,
                "headings": headings,
                "distractors": distractors,
            }
            with open(json_path, "w", encoding="utf-8") as handle:
                json.dump(gt, handle, indent=2, ensure_ascii=False)

            manifest.append({
                "stem": stem,
                "structure": structure_name,
                "variant": variant.name,
                "paragraphs": len(paras),
                "headings": len(headings),
                "distractors": len(distractors),
                "unlocated": unlocated,
            })
            flags = []
            if mismatches:
                flags.append(f"{mismatches} positional mismatch(es)")
            if unlocated:
                flags.append(f"{unlocated} UNLOCATED")
            suffix = f"  [{'; '.join(flags)}]" if flags else ""
            print(
                f"  {stem:34s} {len(paras):3d} paras, {len(headings):2d} headings, "
                f"{len(distractors):2d} distractors{suffix}"
            )

    with open(os.path.join(OUT_DIR, "manifest.json"), "w", encoding="utf-8") as handle:
        json.dump(manifest, handle, indent=2, ensure_ascii=False)

    print(
        f"\n{len(manifest)} documents written; "
        f"{total_unlocated} heading(s) could not be located."
    )
    if total_unlocated:
        print("Unlocated headings carry block_id = -1 and must be excluded from scoring.")


if __name__ == "__main__":
    main()
