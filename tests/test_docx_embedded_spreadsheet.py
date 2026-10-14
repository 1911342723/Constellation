"""DOCX 内嵌电子表格解析（《DOCX 嵌入电子表格解析技术方案》验收）。

夹具全部**合成**：样本文档含账号配置，方案第 5 节明确要求测试夹具用合成或脱敏数据。
所有断言都从 ``DocxProvider.extract_from_bytes`` 真实入口进，不 mock 关系解析或 XLSX
读取——「关系能解析」和「表格真的进了块流」是两件事，只测前者会得到一份永远为空的表。
"""
from __future__ import annotations

import io
import zipfile

import openpyxl
import pytest
from docx import Document

from infrastructure.providers import embedded_spreadsheet as emb
from infrastructure.providers.docx_provider import DocxProvider
from infrastructure.providers.spreadsheet_provider import XlsxProvider

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
O = "urn:schemas-microsoft-com:office:office"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
REL_OLE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject"
CT_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ── 夹具构造 ──────────────────────────────────────────────────────────────


def _workbook(rows, *, title="数据", extra_sheets=None, hidden=False) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for row in rows:
        ws.append(row)
    if hidden:
        ws.sheet_state = "hidden"
    for name, extra_rows in (extra_sheets or []):
        sheet = wb.create_sheet(name)
        for row in extra_rows:
            sheet.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _ole_run_xml(r_id: str, *, prog_id: str = "Excel.Sheet.12", ole_type: str = "Embed") -> str:
    """一个承载 OLE 对象的 w:r（含 VML 预览壳，与 Word 真实产物同形）。"""
    return (
        f'<w:r xmlns:w="{W}" xmlns:r="{R}" xmlns:o="{O}"'
        f' xmlns:v="urn:schemas-microsoft-com:vml">'
        f'<w:object w:dxaOrig="9360" w:dyaOrig="1680">'
        f'<v:shape id="_x0000_i10{r_id[-2:]}" style="width:468pt;height:84pt"/>'
        f'<o:OLEObject Type="{ole_type}" ProgID="{prog_id}" ShapeID="_x0000_i10{r_id[-2:]}"'
        f' DrawAspect="Content" ObjectID="_1{r_id[-3:]}" r:id="{r_id}"/>'
        f'</w:object></w:r>'
    )


def _text_run_xml(text: str) -> str:
    return f'<w:r xmlns:w="{W}"><w:t xml:space="preserve">{text}</w:t></w:r>'


def _build_docx(paragraph_specs, *, embeds, cell_specs=None) -> bytes:
    """按规格造一份带内嵌对象的 DOCX。

    ``paragraph_specs``：每项是 run XML 字符串的列表（一段）。
    ``embeds``：``{rId: (part_name, blob | None, external_target | None)}``。
    ``cell_specs``：可选，往一个 1x1 表格的单元格里放 run XML。
    """
    doc = Document()
    anchors: list = []
    for index, runs in enumerate(paragraph_specs):
        para = doc.add_paragraph()
        anchors.append((para, runs))
    table_para = None
    if cell_specs is not None:
        table = doc.add_table(rows=1, cols=1)
        table_para = table.cell(0, 0).paragraphs[0]

    buf = io.BytesIO()
    doc.save(buf)
    raw = buf.getvalue()

    # 直接改 document.xml：python-docx 没有 OLE 对象 API，注入 XML 是唯一如实造样本的办法。
    with zipfile.ZipFile(io.BytesIO(raw)) as src:
        names = src.namelist()
        doc_xml = src.read("word/document.xml").decode("utf-8")
        rels_xml = src.read("word/_rels/document.xml.rels").decode("utf-8")
        ct_xml = src.read("[Content_Types].xml").decode("utf-8")
        payload = {name: src.read(name) for name in names}

    # 段落：把 python-docx 生成的空段落逐个替换成带 run 的版本。
    empty = "<w:p/>"
    if empty not in doc_xml:
        empty = '<w:p xmlns:w="%s"/>' % W
    for runs in paragraph_specs:
        replacement = "<w:p>" + "".join(runs) + "</w:p>"
        doc_xml = doc_xml.replace(empty, replacement, 1)
    if cell_specs is not None:
        # 表格单元格里的空段落是文档中最后一个 <w:p/>。
        replacement = "<w:p>" + "".join(cell_specs) + "</w:p>"
        doc_xml = doc_xml.replace(empty, replacement, 1)

    rel_entries = []
    for r_id, (part_name, blob, external) in embeds.items():
        if external is not None:
            rel_entries.append(
                f'<Relationship Id="{r_id}" Type="{REL_OLE}" Target="{external}"'
                f' TargetMode="External"/>'
            )
            continue
        rel_entries.append(
            f'<Relationship Id="{r_id}" Type="{REL_OLE}" Target="embeddings/{part_name}"/>'
        )
        if blob is not None:
            payload[f"word/embeddings/{part_name}"] = blob
    rels_xml = rels_xml.replace("</Relationships>", "".join(rel_entries) + "</Relationships>")

    if "embeddings/" in rels_xml and 'Extension="xlsx"' not in ct_xml:
        ct_xml = ct_xml.replace(
            "</Types>", f'<Default Extension="xlsx" ContentType="{CT_XLSX}"/></Types>'
        )
    if 'Extension="bin"' not in ct_xml:
        ct_xml = ct_xml.replace(
            "</Types>",
            '<Default Extension="bin" ContentType="application/vnd.openxmlformats-'
            'officedocument.oleObject"/></Types>',
        )

    payload["word/document.xml"] = doc_xml.encode("utf-8")
    payload["word/_rels/document.xml.rels"] = rels_xml.encode("utf-8")
    payload["[Content_Types].xml"] = ct_xml.encode("utf-8")

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name, data in payload.items():
            dst.writestr(name, data)
    return out.getvalue()


def _tables(blocks):
    return [b for b in blocks if b.type == "table"]


def _embedded(blocks):
    return [b for b in blocks if (b.metadata or {}).get("embedded_spreadsheet")]


# ── P1 识别与读取 ─────────────────────────────────────────────────────────


def test_object_only_paragraph_yields_real_table_block():
    """纯对象段落：内嵌 XLSX 的真实内容进块流，不依赖 Office。"""
    blob = _workbook([["字段", "说明"], ["appId", "应用标识"]])
    docx = _build_docx(
        [[_ole_run_xml("rIdEmb1")]], embeds={"rIdEmb1": ("Sheet1.xlsx", blob, None)}
    )
    blocks = DocxProvider().extract_from_bytes(docx)
    tables = _tables(blocks)
    assert len(tables) == 1
    assert "appId" in tables[0].text and "应用标识" in tables[0].text
    prov = tables[0].metadata["embedded_spreadsheet"]
    assert prov["status"] == emb.STATUS_PARSED
    assert prov["relationship_id"] == "rIdEmb1"
    assert prov["sheet_name"] == "数据"
    assert prov["content_hash"] and len(prov["content_hash"]) == 64


def test_sheet_name_does_not_become_a_heading():
    """工作表名只进 metadata，绝不提升为标题——否则随机表名会篡改章节结构。"""
    blob = _workbook([["a", "b"]], title="Sheet1")
    docx = _build_docx(
        [[_ole_run_xml("rIdEmb1")]], embeds={"rIdEmb1": ("Sheet1.xlsx", blob, None)}
    )
    blocks = DocxProvider().extract_from_bytes(docx)
    assert not [b for b in blocks if b.is_heading_style and (b.text or "").strip() == "Sheet1"]
    assert _tables(blocks)[0].metadata["sheet"] == "Sheet1"


# ── P2 原位集成 ───────────────────────────────────────────────────────────


def test_text_before_and_after_object_keeps_source_order():
    """文字前后夹对象：输出必须是 前文 → 表格 → 后文 三块，顺序与原文一致。"""
    blob = _workbook([["k", "v"]])
    docx = _build_docx(
        [[_text_run_xml("前置说明"), _ole_run_xml("rIdEmb1"), _text_run_xml("后置说明")]],
        embeds={"rIdEmb1": ("Sheet1.xlsx", blob, None)},
    )
    blocks = DocxProvider().extract_from_bytes(docx)
    shape = [(b.type, (b.text or "")[:4]) for b in blocks if (b.text or "").strip()]
    kinds = [t for t, _ in shape]
    assert "table" in kinds
    table_at = kinds.index("table")
    assert shape[table_at - 1][1] == "前置说明"
    assert shape[table_at + 1][1] == "后置说明"


def test_multiple_objects_in_one_paragraph_all_land_in_order():
    blob_a = _workbook([["A1"]])
    blob_b = _workbook([["B1"]])
    docx = _build_docx(
        [[_ole_run_xml("rIdEmbA"), _ole_run_xml("rIdEmbB")]],
        embeds={
            "rIdEmbA": ("A.xlsx", blob_a, None),
            "rIdEmbB": ("B.xlsx", blob_b, None),
        },
    )
    tables = _tables(DocxProvider().extract_from_bytes(docx))
    assert len(tables) == 2
    assert "A1" in tables[0].text and "B1" in tables[1].text


def test_same_attachment_twice_reads_once_but_emits_two_occurrences():
    """同附件重复引用：读取去重（缓存命中一次），正文出现位置**不**去重。"""
    blob = _workbook([["dup", "row"]])
    docx = _build_docx(
        [[_ole_run_xml("rIdEmb1")], [_ole_run_xml("rIdEmb2")]],
        embeds={
            "rIdEmb1": ("Same.xlsx", blob, None),
            "rIdEmb2": ("Same.xlsx", None, None),  # 指向同一部件
        },
    )
    provider = DocxProvider()
    tables = _tables(provider.extract_from_bytes(docx))
    assert len(tables) == 2
    hashes = {t.metadata["embedded_spreadsheet"]["content_hash"] for t in tables}
    occurrences = {t.metadata["embedded_spreadsheet"]["occurrence_id"] for t in tables}
    assert len(hashes) == 1, "同内容应共享 content_hash"
    assert len(occurrences) == 2, "两处出现必须有各自的 occurrence_id"
    assert provider._embedded_cache.hits == 1


def test_occurrence_id_is_stable_across_reparses():
    """occurrence_id 必须由位置确定，重复解析不漂移（不能用随机 UUID）。"""
    blob = _workbook([["x"]])
    docx = _build_docx(
        [[_ole_run_xml("rIdEmb1")]], embeds={"rIdEmb1": ("S.xlsx", blob, None)}
    )
    first = _tables(DocxProvider().extract_from_bytes(docx))[0]
    second = _tables(DocxProvider().extract_from_bytes(docx))[0]
    assert (
        first.metadata["embedded_spreadsheet"]["occurrence_id"]
        == second.metadata["embedded_spreadsheet"]["occurrence_id"]
    )


def test_alternate_content_emits_object_once():
    """mc:AlternateContent 的 Choice/Fallback 是同一对象的两套表示，只能抽一次。"""
    blob = _workbook([["once"]])
    wrapped = (
        f'<mc:AlternateContent xmlns:mc="{MC}">'
        f"<mc:Choice Requires=\"wps\">{_ole_run_xml('rIdEmb1')}</mc:Choice>"
        f"<mc:Fallback>{_ole_run_xml('rIdEmb1')}</mc:Fallback>"
        f"</mc:AlternateContent>"
    )
    docx = _build_docx([[wrapped]], embeds={"rIdEmb1": ("S.xlsx", blob, None)})
    assert len(_tables(DocxProvider().extract_from_bytes(docx))) == 1


def test_object_inside_table_cell_links_back_to_parent_table():
    """单元格内对象：不能无位置地追加，必须带父表 ID 与行列坐标。"""
    blob = _workbook([["cell", "embedded"]])
    docx = _build_docx(
        [[_text_run_xml("正文")]],
        embeds={"rIdEmb1": ("S.xlsx", blob, None)},
        cell_specs=[_ole_run_xml("rIdEmb1")],
    )
    blocks = DocxProvider().extract_from_bytes(docx)
    embedded = _embedded(blocks)
    assert embedded, "单元格内的嵌入表格必须进块流"
    position = embedded[0].metadata["embedded_spreadsheet"]["source_position"]
    assert position.get("parent_table_block_id") is not None
    assert position.get("row") == 0 and position.get("col") == 0


# ── P3 完整性：容错与预算 ─────────────────────────────────────────────────


def test_multi_sheet_reads_all_and_annotates_hidden_state():
    blob = _workbook(
        [["v1"]], title="可见", extra_sheets=[("附表", [["h1"]])]
    )
    tables = _tables(
        DocxProvider().extract_from_bytes(
            _build_docx([[_ole_run_xml("rIdEmb1")]], embeds={"rIdEmb1": ("S.xlsx", blob, None)})
        )
    )
    assert len(tables) == 2
    states = {t.metadata["embedded_spreadsheet"]["sheet_state"] for t in tables}
    assert states == {"visible"}
    assert [t.metadata["embedded_spreadsheet"]["sheet_index"] for t in tables] == [0, 1]


def test_blank_rows_keep_original_coordinates():
    """裁掉空行后必须保留原始行号，否则任何按坐标回定位都会错位。"""
    blob = _workbook([["h1", "h2"], [None, None], ["r3a", "r3b"]])
    table = _tables(
        DocxProvider().extract_from_bytes(
            _build_docx([[_ole_run_xml("rIdEmb1")]], embeds={"rIdEmb1": ("S.xlsx", blob, None)})
        )
    )[0]
    prov = table.metadata["embedded_spreadsheet"]
    assert prov["row_indices"] == [0, 2]
    assert prov["read_range"]["rows"] == 2


def test_formula_without_cached_value_is_not_treated_as_blank():
    """无缓存值的公式格不得变成空白——那是静默丢数据。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["数量", "合计"])
    ws.append([3, "=A2*2"])
    buf = io.BytesIO()
    wb.save(buf)
    docx = _build_docx(
        [[_ole_run_xml("rIdEmb1")]], embeds={"rIdEmb1": ("S.xlsx", buf.getvalue(), None)}
    )
    provider = DocxProvider()
    tables = _tables(provider.extract_from_bytes(docx))
    assert "=A2*2" in tables[0].text
    codes = {skip.code for skip in provider.drain_skips()}
    assert "xlsx_formula_no_cached_value" in codes


def test_missing_relationship_leaves_visible_placeholder():
    docx = _build_docx([[_ole_run_xml("rIdNope")]], embeds={})
    provider = DocxProvider()
    blocks = provider.extract_from_bytes(docx)
    placeholder = [b for b in blocks if "嵌入电子表格未解析" in (b.text or "")]
    assert placeholder, "关系缺失必须留正文可见占位，不能静默跳过"
    assert placeholder[0].metadata["embedded_spreadsheet"]["status"] == emb.STATUS_FAILED
    assert emb.WARN_MISSING in {skip.code for skip in provider.drain_skips()}


def test_external_link_object_is_registered_without_network():
    docx = _build_docx(
        [[_ole_run_xml("rIdExt", ole_type="Link")]],
        embeds={"rIdExt": (None, None, "https://example.invalid/book.xlsx")},
    )
    provider = DocxProvider()
    blocks = provider.extract_from_bytes(docx)
    embedded = _embedded(blocks)
    assert embedded[0].metadata["embedded_spreadsheet"]["status"] == emb.STATUS_EXTERNAL
    assert emb.WARN_EXTERNAL in {skip.code for skip in provider.drain_skips()}


def test_ole_compound_binary_is_reported_unsupported():
    """OLE 复合文档（旧版 XLS 容器）本期不支持，如实标注而不是喂给 openpyxl 炸掉。"""
    docx = _build_docx(
        [[_ole_run_xml("rIdBin", prog_id="Excel.Sheet.8")]],
        embeds={"rIdBin": ("oleObject1.bin", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64, None)},
    )
    provider = DocxProvider()
    blocks = provider.extract_from_bytes(docx)
    prov = _embedded(blocks)[0].metadata["embedded_spreadsheet"]
    assert prov["status"] == emb.STATUS_UNSUPPORTED
    assert emb.WARN_UNSUPPORTED in {skip.code for skip in provider.drain_skips()}


def test_corrupt_package_fails_only_that_object():
    """损坏附件只让该对象失败，正文其余内容照常输出。"""
    good = _workbook([["ok", "row"]])
    docx = _build_docx(
        [[_ole_run_xml("rIdBad")], [_text_run_xml("正文继续")], [_ole_run_xml("rIdGood")]],
        embeds={
            "rIdBad": ("bad.xlsx", b"PK\x03\x04broken-not-a-zip", None),
            "rIdGood": ("good.xlsx", good, None),
        },
    )
    provider = DocxProvider()
    blocks = provider.extract_from_bytes(docx)
    assert any("正文继续" in (b.text or "") for b in blocks)
    assert any("ok" in (b.text or "") for b in _tables(blocks))
    assert any("嵌入电子表格未解析" in (b.text or "") for b in blocks)


def test_document_level_budget_is_shared_not_reset_per_workbook():
    """预算是文档级共享的：两个附件各自合法但合计越限时，第二个必须被挡住。"""
    blob_a = _workbook([[f"a{i}"] for i in range(12)])
    blob_b = _workbook([[f"b{i}"] for i in range(12)])
    docx = _build_docx(
        [[_ole_run_xml("rIdA")], [_ole_run_xml("rIdB")]],
        embeds={"rIdA": ("A.xlsx", blob_a, None), "rIdB": ("B.xlsx", blob_b, None)},
    )
    provider = DocxProvider(embedded_budget_factory=lambda: emb.EmbeddedBudget(max_rows=15))
    blocks = provider.extract_from_bytes(docx)
    tables = _tables(blocks)
    assert tables, "预算内的第一个附件仍应解析"
    assert provider._embedded_budget.rows_read <= 15, "累计行数不得越过文档级预算"
    # 关键：第二个附件拿到的是**剩余**额度而不是重置后的满额度，所以它必须被截断。
    assert tables[-1].metadata["embedded_spreadsheet"]["read_truncated"] is True
    assert tables[-1].metadata["embedded_spreadsheet"]["status"] == emb.STATUS_PARTIAL


def test_object_count_over_budget_leaves_placeholder_and_warns():
    blob = _workbook([["x"]])
    docx = _build_docx(
        [[_ole_run_xml("rIdA")], [_ole_run_xml("rIdB")]],
        embeds={"rIdA": ("A.xlsx", blob, None), "rIdB": ("B.xlsx", blob, None)},
    )
    provider = DocxProvider(embedded_budget_factory=lambda: emb.EmbeddedBudget(max_objects=1))
    blocks = provider.extract_from_bytes(docx)
    assert emb.WARN_BUDGET in {skip.code for skip in provider.drain_skips()}
    assert any("超出文档级预算" in (b.text or "") for b in blocks)


# ── P4 回归：无嵌入对象的文档行为不变 ──────────────────────────────────────


def test_object_preview_is_not_emitted_as_separate_content():
    """表格解析成功后，对象的 VML 预览图不再作为独立语义内容输出（同一来源不重复）。"""
    blob = _workbook([["only", "once"]])
    docx = _build_docx(
        [[_ole_run_xml("rIdEmb1")]], embeds={"rIdEmb1": ("S.xlsx", blob, None)}
    )
    blocks = DocxProvider().extract_from_bytes(docx)
    assert not [b for b in blocks if b.type == "image"], "预览图不得与表格重复输出"
    assert len(_tables(blocks)) == 1


def test_provenance_reaches_document_ir_not_only_provider():
    """来源必须能被下游取回——方案第 4 节禁止「仅在 Provider 内可见」。"""
    from modules.parser.document_ir import NodeKind, ir_from_block

    blob = _workbook([["ir", "passthrough"]])
    docx = _build_docx(
        [[_ole_run_xml("rIdEmb1")]], embeds={"rIdEmb1": ("S.xlsx", blob, None)}
    )
    table = _tables(DocxProvider().extract_from_bytes(docx))[0]
    node = ir_from_block(table)
    assert node.kind == NodeKind.TABLE
    prov = node.source.provenance.get("embedded_spreadsheet")
    assert prov is not None, "IR 丢掉了嵌入对象来源"
    assert prov["status"] == emb.STATUS_PARSED
    assert prov["occurrence_id"] == table.metadata["embedded_spreadsheet"]["occurrence_id"]
    assert node.table.rows and "passthrough" in node.table.rows[0]


def test_source_range_reports_unknown_instead_of_a_wrong_upper_bound():
    """真实样本暴露的缺陷：内嵌工作簿常**不写 dimension**，只读探针只能回 1。

    把 1 当"原始行数上界"上报，就会出现「上界 1 < 实际读到 33 行」这种自相矛盾的数字，
    下游据此判断"有没有被截断"会得到反的结论。所以估不出来必须标 known=False。
    """
    blob = _workbook([[f"r{i}", f"v{i}"] for i in range(8)])
    table = _tables(
        DocxProvider().extract_from_bytes(
            _build_docx([[_ole_run_xml("rIdEmb1")]], embeds={"rIdEmb1": ("S.xlsx", blob, None)})
        )
    )[0]
    source_range = table.metadata["embedded_spreadsheet"]["source_range"]
    read_rows = table.metadata["embedded_spreadsheet"]["read_range"]["rows"]
    assert read_rows == 8
    if source_range.get("known"):
        # 估得出来时必须是真正的上界，不能低于实际已读。
        assert source_range["rows"] >= read_rows
    else:
        assert source_range["rows"] is None
        assert "dimension" in source_range["unknown_reason"]


def test_plain_docx_without_objects_is_unchanged():
    doc = Document()
    doc.add_heading("标题", level=1)
    doc.add_paragraph("普通段落")
    buf = io.BytesIO()
    doc.save(buf)
    blocks = DocxProvider().extract_from_bytes(buf.getvalue())
    assert [b.id for b in blocks] == list(range(len(blocks)))
    assert not _embedded(blocks)
    assert any((b.text or "").strip() == "普通段落" for b in blocks)


def test_block_ids_stay_contiguous_with_embedded_objects():
    blob = _workbook([["a"]], extra_sheets=[("二", [["b"]])])
    docx = _build_docx(
        [[_text_run_xml("前")], [_ole_run_xml("rIdEmb1")], [_text_run_xml("后")]],
        embeds={"rIdEmb1": ("S.xlsx", blob, None)},
    )
    blocks = DocxProvider().extract_from_bytes(docx)
    assert [b.id for b in blocks] == list(range(len(blocks)))


# ── 方案 §8「必要自动化覆盖」补齐的六项 ────────────────────────────────────


def test_empty_embedded_sheet_yields_no_content_free_table():
    """§8③ 空工作表：不能装出一个**没有任何内容**的表格块。

    嵌入模式没有工作表名标题，所以一个空表块在正文里就是纯噪声；独立 XLSX 路径早有
    ``if not rows: continue`` 守卫，嵌入路径此前漏了。全空时落「工作簿无可读内容」占位。
    """
    blob = _workbook([])  # 一张什么都没有的工作表
    docx = _build_docx(
        [[_ole_run_xml("rIdEmpty")]], embeds={"rIdEmpty": ("Empty.xlsx", blob, None)}
    )
    provider = DocxProvider()
    blocks = provider.extract_from_bytes(docx)
    assert not _tables(blocks), "空工作表不得产出表格块"
    placeholder = _embedded(blocks)
    assert placeholder, "全空的工作簿仍须在原位留可见占位"
    assert placeholder[0].metadata["embedded_spreadsheet"]["status"] == emb.STATUS_PARTIAL
    assert "xlsx_sheet_empty" in {skip.code for skip in provider.drain_skips()}


def test_empty_sheet_alongside_a_real_one_keeps_the_real_one():
    """一空一实：空的那张不产块，实的那张照常出——空表不能连坐。"""
    blob = _workbook([["有", "数据"]], title="实表", extra_sheets=[("空表", [])])
    docx = _build_docx(
        [[_ole_run_xml("rIdMix")]], embeds={"rIdMix": ("Mix.xlsx", blob, None)}
    )
    provider = DocxProvider()
    tables = _tables(provider.extract_from_bytes(docx))
    assert len(tables) == 1
    assert tables[0].metadata["embedded_spreadsheet"]["sheet_name"] == "实表"
    assert "xlsx_sheet_empty" in {skip.code for skip in provider.drain_skips()}


def test_merged_cells_in_embedded_workbook_are_filled_down():
    """§8③ 合并单元格：代码路径有 ``scan_merged_cells``，此前从未在嵌入场景验证。

    合并区只有左上角存值，其余读回 None。不做 fill-down 的表现是整张表右移错列——
    渲染出来仍然「有文字」，只有逐格断言抓得住。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "合并"
    ws["A1"] = "季度汇总"
    ws.merge_cells("A1:C1")
    ws.append(["一月", "二月", "三月"])
    ws.append([1, 2, 3])
    buf = io.BytesIO()
    wb.save(buf)
    docx = _build_docx(
        [[_ole_run_xml("rIdMerge")]], embeds={"rIdMerge": ("M.xlsx", buf.getvalue(), None)}
    )
    table = _tables(DocxProvider().extract_from_bytes(docx))[0]
    rows = table.table_data["rows"]
    assert rows[0] == ["季度汇总", "季度汇总", "季度汇总"], "合并区未 fill-down，后续列会错位"
    assert rows[1] == ["一月", "二月", "三月"]


def test_merged_cells_survive_the_streaming_path_too():
    """嵌入场景有两条读取路径，合并语义在**两条**上都必须成立。

    小表走 ``read_only=False``（openpyxl 直接给合并区），大表走流式 ``read_only=True``
    （合并信息在该模式不可靠，靠 ``scan_merged_cells`` 独立扫 XML 补）。剩余行预算低于
    ``max_rows`` 时会切到流式，这里就用文档级预算把它逼过去——否则「流式 + 合并」这条
    分支永远跑不到，而它恰恰是自己扫 XML 的那一条。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "流式合并"
    ws["A1"] = "跨列表头"
    ws.merge_cells("A1:B1")
    for i in range(6):
        ws.append([f"r{i}a", f"r{i}b"])
    buf = io.BytesIO()
    wb.save(buf)
    docx = _build_docx(
        [[_ole_run_xml("rIdStream")]], embeds={"rIdStream": ("S.xlsx", buf.getvalue(), None)}
    )
    # max_rows 远小于 XlsxProvider 的 5000 → sheet_row_cap < max_rows → is_large → 流式。
    provider = DocxProvider(embedded_budget_factory=lambda: emb.EmbeddedBudget(max_rows=50))
    table = _tables(provider.extract_from_bytes(docx))[0]
    assert table.table_data["rows"][0] == ["跨列表头", "跨列表头"]


def _workbook_with_cached_formula() -> bytes:
    """造一个公式**带缓存值**的工作簿。

    openpyxl 写不出缓存值（它只写 ``<f>``），所以保存后直接改 sheet XML 塞 ``<v>``——
    这正是 Excel 真实产物的形态。不这么造，「有缓存」那一侧就永远测不到。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "有缓存"
    ws.append(["数量", "合计"])
    ws.append([3, "=A2*2"])
    buf = io.BytesIO()
    wb.save(buf)

    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as src:
        payload = {name: src.read(name) for name in src.namelist()}
    sheet_name = next(n for n in payload if n.startswith("xl/worksheets/sheet"))
    xml = payload[sheet_name].decode("utf-8")
    assert "<f>A2*2</f>" in xml, "夹具前提变了：openpyxl 不再这样写公式"
    xml = xml.replace("<f>A2*2</f>", "<f>A2*2</f><v>6</v>")
    payload[sheet_name] = xml.encode("utf-8")

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name, data in payload.items():
            dst.writestr(name, data)
    return out.getvalue()


def test_formula_with_cached_value_uses_the_value_and_does_not_warn():
    """§8④ 有缓存值的公式：此前只测了「无缓存」那一侧。

    有缓存时必须用缓存值（``6``），且**不得**记 ``xlsx_formula_no_cached_value``——
    那条告警的含义是「这里给的是公式源文本而非计算值」，缓存齐全时报它就是假警报，
    会让整份结果无端变成 partial。
    """
    docx = _build_docx(
        [[_ole_run_xml("rIdCached")]],
        embeds={"rIdCached": ("C.xlsx", _workbook_with_cached_formula(), None)},
    )
    provider = DocxProvider()
    table = _tables(provider.extract_from_bytes(docx))[0]
    rows = table.table_data["rows"]
    assert rows[1] == ["3", "6"], f"应使用缓存计算值，实际 {rows[1]}"
    assert "=A2*2" not in table.text
    assert "xlsx_formula_no_cached_value" not in {skip.code for skip in provider.drain_skips()}


def _zip_bomb_workbook(entry_count: int) -> bytes:
    """在一个合法 XLSX 包里塞大量额外条目——触发「单包条目数」这道门。"""
    base = _workbook([["x"]])
    with zipfile.ZipFile(io.BytesIO(base)) as src:
        payload = {name: src.read(name) for name in src.namelist()}
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name, data in payload.items():
            dst.writestr(name, data)
        for index in range(entry_count):
            dst.writestr(f"xl/filler/pad{index}.bin", b"0" * 32)
    return out.getvalue()


def test_inner_package_entry_explosion_is_refused_before_reading():
    """§8⑥ 内层解压超限：``claim_package`` 此前有实现、无用例。

    关键是「读之前」就拒——门开在 ``sniff`` 之后、``read_embedded`` 之前，靠中央目录
    的条目数与展开体积判断，不解压。所以断言不仅看告警，还看**一行都没读**。
    """
    docx = _build_docx(
        [[_ole_run_xml("rIdBomb")]],
        embeds={"rIdBomb": ("Bomb.xlsx", _zip_bomb_workbook(64), None)},
    )
    provider = DocxProvider(
        embedded_budget_factory=lambda: emb.EmbeddedBudget(max_entries_per_package=8)
    )
    blocks = provider.extract_from_bytes(docx)
    assert not _tables(blocks), "越过条目数门的包不得被读取"
    prov = _embedded(blocks)[0].metadata["embedded_spreadsheet"]
    assert prov["status"] == emb.STATUS_FAILED
    assert "条目数" in prov["detail"]
    assert emb.WARN_BUDGET in {skip.code for skip in provider.drain_skips()}
    assert provider._embedded_budget.rows_read == 0, "应在读取前拒绝，不该已经读进行"


def test_inner_package_expanded_bytes_ceiling_is_cumulative():
    """展开体积门是**累计**的：第一个包吃掉额度后，第二个必须被挡在读取前。"""
    blob_a = _workbook([[f"a{i}", "x" * 200] for i in range(30)])
    blob_b = _workbook([[f"b{i}", "y" * 200] for i in range(30)])
    docx = _build_docx(
        [[_ole_run_xml("rIdA")], [_ole_run_xml("rIdB")]],
        embeds={"rIdA": ("A.xlsx", blob_a, None), "rIdB": ("B.xlsx", blob_b, None)},
    )
    # 阈值取「一个包够、两个包不够」：用第一个包的实际展开体积做基准。
    first = emb.sniff_embedded_format(blob_a)
    assert first.expanded_bytes > 0
    ceiling = first.expanded_bytes + max(first.expanded_bytes // 2, 1)
    provider = DocxProvider(
        embedded_budget_factory=lambda: emb.EmbeddedBudget(max_expanded_bytes=ceiling)
    )
    blocks = provider.extract_from_bytes(docx)
    tables = _tables(blocks)
    assert tables, "额度内的第一个包仍应解析"
    assert all("b0" not in (t.text or "") for t in tables), "第二个包不该被读到"
    refused = [
        b for b in _embedded(blocks)
        if b.metadata["embedded_spreadsheet"]["status"] == emb.STATUS_FAILED
    ]
    assert refused and "展开体积" in refused[0].metadata["embedded_spreadsheet"]["detail"]


def test_many_attachments_share_one_read_budget_across_the_document():
    """§8⑥ 多附件压力夹具：10 个附件共享一份文档级读取预算。

    断言落在 ``read_range``（**真实已读**行数）与预算账本上，而不是渲染出来的文本——
    渲染层本来就会把大表抽样，「输出看起来很小」在**已经把全部行读进内存**的情况下
    同样成立，只看输出等于没测。
    """
    count = 10
    rows_per_book = 60
    embeds = {}
    paragraphs = []
    for index in range(count):
        blob = _workbook(
            [[f"s{index}r{row}", "填充" * 20] for row in range(rows_per_book)],
            title=f"表{index}",
        )
        r_id = f"rIdBulk{index}"
        embeds[r_id] = (f"Bulk{index}.xlsx", blob, None)
        paragraphs.append([_ole_run_xml(r_id)])
    docx = _build_docx(paragraphs, embeds=embeds)

    provider = DocxProvider(embedded_budget_factory=lambda: emb.EmbeddedBudget(max_rows=80))
    blocks = provider.extract_from_bytes(docx)
    tables = _tables(blocks)
    read_rows = sum(
        t.metadata["embedded_spreadsheet"]["read_range"]["rows"] for t in tables
    )

    assert tables, "预算内的前几个附件仍应解析"
    assert read_rows <= 80, f"累计真实已读 {read_rows} 越过了文档级预算"
    assert provider._embedded_budget.rows_read <= 80
    # 预算耗尽后剩下的附件必须留可见占位 / 截断标记，而不是静默消失。
    assert len(tables) < count
    assert any(
        t.metadata["embedded_spreadsheet"]["status"] in (emb.STATUS_PARTIAL,)
        for t in tables
    )


def test_read_budget_bounds_peak_allocation_not_only_rendered_text():
    """§8「性能记录」：预算限制的是累计**读取**，所以峰值分配必须随预算下降。

    量在 ``read_embedded`` 这一层而不是整份 DOCX 上：整份解析的峰值被夹具字节本身
    主导（两种预算下完全相同），差异会被噪声淹没，那样的断言是空过的。

    夹具用**数值**单元格：read_only 模式会把共享字符串表整体载入，若用字符串填充，
    两种预算的峰值都等于字符串表大小，差异同样消失——数值内联在 sheet XML 里，
    才真正体现「读多少行就分配多少」。

    tracemalloc 量的是 Python 分配字节，不是进程 RSS（Windows 上 ``resource`` 模块
    不存在，指标面板对峰值 RSS 如实标「本环境不可测」）。这是如实口径，不拿别的数字冒充。
    """
    import tracemalloc

    total_rows = 3000
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "压力"
    for row in range(total_rows):
        ws.append([row * 10 + col for col in range(5)])
    buf = io.BytesIO()
    wb.save(buf)
    blob = buf.getvalue()

    def read(max_rows: int) -> tuple[int, int]:
        provider = XlsxProvider()
        budget = emb.EmbeddedBudget(max_rows=max_rows)
        tracemalloc.start()
        try:
            sheets = provider.read_embedded(blob, budget=budget)
            _current, peak = tracemalloc.get_traced_memory()
        finally:
            tracemalloc.stop()
        provider.drain_skips()
        return sum(len(s.rows) for s in sheets), peak

    tight_rows, tight_peak = read(50)
    loose_rows, loose_peak = read(total_rows)

    assert tight_rows == 50 and loose_rows == total_rows
    # 留足倍数余量：只要求「显著下降」，不钉一个会随 openpyxl 版本漂移的绝对值。
    assert tight_peak * 2 < loose_peak, (
        f"紧预算峰值 {tight_peak} 未显著低于宽预算峰值 {loose_peak}——"
        f"说明预算只管住了渲染，没管住读取"
    )


def test_render_sampling_and_read_truncation_are_different_things():
    """预算与渲染抽样是两件事，metadata 必须分得开（方案第 5 节点名的一条）。

    一张读全了但只渲染前若干行的表：``read_truncated=False`` + ``render_sampled=True``。
    混成一个字段，partial 语义就失真——「显示抽样」会被当成「数据不完整」。
    """
    rows = 260  # > max_render_rows(200) 但远小于 max_rows(5000)
    blob = _workbook([[f"r{i}", i] for i in range(rows)])
    docx = _build_docx(
        [[_ole_run_xml("rIdSample")]], embeds={"rIdSample": ("S.xlsx", blob, None)}
    )
    table = _tables(DocxProvider().extract_from_bytes(docx))[0]
    prov = table.metadata["embedded_spreadsheet"]
    assert prov["read_range"]["rows"] == rows, "数据应被完整读入"
    assert prov["read_truncated"] is False
    assert prov["render_sampled"] is True
    assert prov["status"] == emb.STATUS_PARSED
    # table_data 是真实已读集合，不受渲染抽样影响。
    assert len(table.table_data["rows"]) == rows
    assert table.text.count("\n| ") < rows


def test_header_footer_parts_are_declared_unsupported_not_silently_dropped():
    """§8② 的「不同 part 的同名关系 ID」**无法覆盖**：首期只做 document part。

    页眉页脚 part 明确不在范围内（方案 §9「未做」列了这一条）。这条用例不去实现它，
    只把「不支持」这个声明钉住：一旦哪天有人顺手扫了页眉页脚，声明与实现就会不一致，
    这里会红——比在文档里写一句「暂不支持」有效。
    """
    from infrastructure.providers import docx_provider as dp

    source = (
        __import__("pathlib").Path(dp.__file__).read_text(encoding="utf-8")
    )
    assert "word/header" not in source and "word/footer" not in source, (
        "出现了页眉页脚 part 的处理：请同步更新方案 §9「未做」清单与本用例"
    )
