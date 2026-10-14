"""B01/B04 验收：DOCX 包装节点文字丢失、纵向合并 row_span 不准。

这些用例是方案 B01/B04 的"最小失败样例"：修复前它们全部失败
（超链接/插入文字被静默吞掉、合并格文字逐行重复且 row_span 恒为 1），
用来钉死回归。构造真实 OOXML 片段而非依赖 python-docx 的高层 API，
因为缺陷正出在 python-docx 对 w:hyperlink/w:ins/vMerge 的展平处理上。
"""
from io import BytesIO
import sys
import zipfile
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).parent.parent))

from infrastructure.providers.docx_provider import DocxProvider


def _save(doc: Document) -> bytes:
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _rewrite(file_bytes: bytes, entry: str, old: bytes, new: bytes) -> bytes:
    """在 document.xml 里替换标记片段，注入手工构造的 OOXML 结构。"""
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(file_bytes)) as zin, zipfile.ZipFile(output, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == entry:
                assert old in data, f"marker {old!r} not found in {entry}"
                data = data.replace(old, new)
            zout.writestr(info, data)
    return output.getvalue()


def _para_with(marker: str, injected: bytes) -> bytes:
    """把一个含 marker 文字的普通段落，改造成注入 OOXML 片段的段落。"""
    doc = Document()
    doc.add_paragraph(marker)
    return _rewrite(
        _save(doc),
        "word/document.xml",
        f"<w:t>{marker}</w:t>".encode("utf-8"),
        injected,
    )


def _para_replacing_run(marker: str, injected: bytes) -> bytes:
    """把整个 `<w:r><w:t>marker</w:t></w:r>` 替换为注入片段。

    包装节点（hyperlink / ins / sdt）是段落级子元素，必须替换整个 run，
    否则会被塞进 w:r 内部形成非法 OOXML（迭代器只认段落直属子节点）。
    """
    doc = Document()
    doc.add_paragraph(marker)
    return _rewrite(
        _save(doc),
        "word/document.xml",
        f"<w:r><w:t>{marker}</w:t></w:r>".encode("utf-8"),
        injected,
    )


def _first_text_block(file_bytes: bytes):
    blocks = DocxProvider().extract_from_bytes(file_bytes)
    return next(b for b in blocks if b.type == "text")


def _only_table(file_bytes: bytes):
    blocks = DocxProvider().extract_from_bytes(file_bytes)
    return next(b for b in blocks if b.type == "table")


# ============================ B01 ============================

_HYPERLINK = (
    b'<w:hyperlink w:anchor="target1"><w:r><w:t>LINKTEXT</w:t></w:r></w:hyperlink>'
)
_INS = b'<w:ins w:id="7" w:author="a"><w:r><w:t>INSERTED</w:t></w:r></w:ins>'
_DEL = b'<w:del w:id="8" w:author="a"><w:r><w:delText>DELETED</w:delText></w:r></w:del>'


def test_b01_internal_hyperlink_anchor_text_survives_in_mixed_paragraph():
    file_bytes = _para_with(
        "PREFIX MID SUFFIX",
        b'<w:t xml:space="preserve">PREFIX </w:t></w:r>'
        + _HYPERLINK
        + b'<w:r><w:t xml:space="preserve"> SUFFIX</w:t>',
    )
    block = _first_text_block(file_bytes)
    assert block.text == "PREFIX [LINKTEXT](#target1) SUFFIX"


def test_b01_hyperlink_only_paragraph_is_not_dropped_by_emptiness_check():
    file_bytes = _para_replacing_run("X", _HYPERLINK)
    block = _first_text_block(file_bytes)
    assert block.text == "[LINKTEXT](#target1)"


def test_b01_inserted_revision_text_is_kept_in_accepted_view():
    file_bytes = _para_with(
        "BEFORE MID AFTER",
        b'<w:t xml:space="preserve">BEFORE </w:t></w:r>'
        + _INS
        + b'<w:r><w:t xml:space="preserve"> AFTER</w:t>',
    )
    block = _first_text_block(file_bytes)
    assert "INSERTED" in block.text
    assert block.text.startswith("BEFORE ") and block.text.endswith(" AFTER")
    assert block.metadata["revision_view"] == "accepted"
    assert {"type": "ins", "text": "INSERTED"} in block.metadata["revisions"]


def test_b01_inserted_only_paragraph_is_not_dropped():
    file_bytes = _para_replacing_run("X", _INS)
    block = _first_text_block(file_bytes)
    assert block.text == "INSERTED"


def test_b01_deleted_text_is_removed_in_accepted_view_but_auditable():
    file_bytes = _para_with(
        "KEEP MID END",
        b'<w:t xml:space="preserve">KEEP</w:t></w:r>'
        + _DEL
        + b'<w:r><w:t xml:space="preserve">END</w:t>',
    )
    block = _first_text_block(file_bytes)
    assert "DELETED" not in block.text
    assert block.text == "KEEPEND"
    # 被删除内容不进正文，但逐条记录可审计（markup 视图可重建）。
    assert {"type": "del", "text": "DELETED"} in block.metadata["revisions"]


def test_b01_content_control_sdt_text_survives():
    sdt = (
        b'<w:sdt><w:sdtContent><w:r><w:t>SDTTEXT</w:t></w:r></w:sdtContent></w:sdt>'
    )
    file_bytes = _para_with(
        "PRE MID POST",
        b'<w:t xml:space="preserve">PRE </w:t></w:r>'
        + sdt
        + b'<w:r><w:t xml:space="preserve"> POST</w:t>',
    )
    block = _first_text_block(file_bytes)
    assert block.text == "PRE SDTTEXT POST"


def test_b01_nested_ins_inside_hyperlink_survives():
    nested = (
        b'<w:hyperlink w:anchor="n1">'
        b'<w:ins w:id="9" w:author="a"><w:r><w:t>NESTED</w:t></w:r></w:ins>'
        b'</w:hyperlink>'
    )
    file_bytes = _para_replacing_run("X", nested)
    block = _first_text_block(file_bytes)
    assert block.text == "[NESTED](#n1)"
    assert {"type": "ins", "text": "NESTED"} in block.metadata["revisions"]


def test_b01_hyperlink_run_becomes_auditable_atom():
    file_bytes = _para_with(
        "A MID B",
        b'<w:t xml:space="preserve">A </w:t></w:r>'
        + _HYPERLINK
        + b'<w:r><w:t xml:space="preserve"> B</w:t>',
    )
    block = _first_text_block(file_bytes)
    atom_texts = [atom["text"] for atom in block.metadata["atoms"]]
    # 来源顺序可还原：三段（含超链接锚文本）都各自成 atom，顺序与阅读一致。
    assert atom_texts == ["A ", "LINKTEXT", " B"]


def test_b01_hyperlink_inside_table_cell_is_not_dropped():
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "CELLMARK"
    file_bytes = _rewrite(
        _save(doc),
        "word/document.xml",
        b"<w:t>CELLMARK</w:t>",
        b'<w:t xml:space="preserve">CELL </w:t></w:r>' + _HYPERLINK + b"<w:r>",
    )
    table_block = _only_table(file_bytes)
    assert "[LINKTEXT](#target1)" in table_block.table_data["rows"][0][0]


# ============================ B04 ============================


def test_b04_vertical_merge_reports_true_row_span_without_duplicating_text():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "MERGED"
    table.cell(0, 1).text = "R0C1"
    table.cell(1, 1).text = "R1C1"
    table.cell(0, 0).merge(table.cell(1, 0))

    table_block = _only_table(_save(doc))
    cells = table_block.metadata["cells"]
    origin = next(c for c in cells if c["text"] == "MERGED")

    assert origin["row_span"] == 2
    assert origin["col_span"] == 1
    assert {"row": 1, "col": 0} in origin["covered"]
    # 覆盖位置不得复制起始格文字冒充新增内容（GFM 降级后为空）。
    assert table_block.table_data["rows"][1][0] == ""
    assert sum(row.count("MERGED") for row in table_block.table_data["rows"]) == 1
    assert table_block.metadata["gfm_downgrade"]["has_vertical_merge"] is True
    assert table_block.metadata["gfm_downgrade"]["spans_representable"] is False


def test_b04_horizontal_merge_reports_col_span():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "WIDE"
    table.cell(1, 0).text = "a"
    table.cell(1, 1).text = "b"
    table.cell(0, 0).merge(table.cell(0, 1))

    table_block = _only_table(_save(doc))
    origin = next(c for c in table_block.metadata["cells"] if c["text"] == "WIDE")
    assert origin["col_span"] == 2
    assert origin["row_span"] == 1


def test_b04_rectangular_merge_reports_both_spans():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "BLOCK"
    # 对角合并即得 2x2 矩形合并格（python-docx 要求跨度矩形）。
    merged = table.cell(0, 0).merge(table.cell(1, 1))
    merged.text = "BLOCK"

    table_block = _only_table(_save(doc))
    origin = next(c for c in table_block.metadata["cells"] if c["text"] == "BLOCK")
    assert origin["col_span"] == 2
    assert origin["row_span"] == 2
    # 覆盖 (0,1) / (1,0) / (1,1) 三格；起始格自身不计入。
    assert len(origin["covered"]) == 3


def test_b04_multi_paragraph_merged_cell_keeps_all_content_once():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    cell = table.cell(0, 0)
    cell.text = "LINE1"
    cell.add_paragraph("LINE2")
    table.cell(0, 1).text = "x"
    table.cell(1, 1).text = "y"
    table.cell(0, 0).merge(table.cell(1, 0))

    table_block = _only_table(_save(doc))
    origin = next(c for c in table_block.metadata["cells"] if "LINE1" in c["text"])
    assert origin["row_span"] == 2
    assert "LINE1" in origin["text"] and "LINE2" in origin["text"]
    # 多段内容也只出现一次，不因合并逐行重复。
    joined = "\n".join("".join(row) for row in table_block.table_data["rows"])
    assert joined.count("LINE1") == 1


def test_b04_nested_table_does_not_leak_inner_rows_into_outer_grid():
    """嵌套表：内层 w:tr 不得被算进外层网格（否则外层跨度全错）。

    grid 解析用 findall('w:tr') 只取直属行，天然隔离内层表；这里钉死该隔离。
    E05(3) 之后内层表**会**单独抽成 Block 并与外层格双向建 span 关系（见
    tests/test_docx_numbering_and_notes.py::test_e05_nested_table_becomes_its_own_block_...），
    但本用例仍只验证"外层结构不被内层污染"这一独立性质：外层 1x1 网格、
    单元格跨度不因内层 2 行而虚报。
    """
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=2, cols=2)
    inner.cell(0, 0).text = "IN"
    inner.cell(0, 1).text = "p"
    inner.cell(1, 1).text = "q"
    inner.cell(0, 0).merge(inner.cell(1, 0))

    blocks = DocxProvider().extract_from_bytes(_save(doc))
    tables = [b for b in blocks if b.type == "table"]
    assert tables, "外层表必须被提取"
    outer_block = tables[0]
    # 外层是 1x1，不能因内层 2 行而虚报成 2 行/2 列。
    assert outer_block.metadata["grid_size"] == {"rows": 1, "cols": 1}
    assert len(outer_block.metadata["cells"]) == 1
    assert outer_block.metadata["cells"][0]["row_span"] == 1
    assert outer_block.metadata["cells"][0]["col_span"] == 1
