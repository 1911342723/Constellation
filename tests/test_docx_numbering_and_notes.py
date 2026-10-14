"""E05 验收：DOCX 编号 / 脚注尾注 / 书签内部引用 / cell 内嵌内容 / 嵌套表入块 /
OMML 转换质量标记 / SmartArt·OLE·未知 XML 受控占位 / P02 code 合并逐字符等价。

全部用例从 ``DocxProvider.extract_from_bytes`` 真实入口进——要么用 python-docx
的高层 API 造样例，要么按 ``tests/test_docx_inline_and_merge.py`` 的做法直接改写
document.xml 注入手工 OOXML（编号定义 / startOverride / 脚注 / SmartArt 等
python-docx 不便构造的结构）。不写"mock 掉调用方后本函数行为正确"的用例。
"""
from io import BytesIO
import sys
import zipfile
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).parent.parent))

from infrastructure.providers.docx_provider import DocxProvider


# ---------------------------------------------------------------------------
# 造样例的工具：改写 docx 内某个 zip entry / 追加新 part + 关系
# ---------------------------------------------------------------------------

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _save(doc: Document) -> bytes:
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _blocks(file_bytes: bytes):
    return DocxProvider().extract_from_bytes(file_bytes)


def _rewrite(file_bytes: bytes, entry: str, old: bytes, new: bytes) -> bytes:
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(file_bytes)) as zin, zipfile.ZipFile(output, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == entry:
                assert old in data, f"marker {old!r} not found in {entry}"
                data = data.replace(old, new)
            zout.writestr(info, data)
    return output.getvalue()


def _add_parts(file_bytes: bytes, parts: dict, rels: list, content_types: list) -> bytes:
    """向 docx 追加新 part、document.xml.rels 关系、[Content_Types].xml Override。"""
    output = BytesIO()
    rels_path = "word/_rels/document.xml.rels"
    ct_path = "[Content_Types].xml"
    with zipfile.ZipFile(BytesIO(file_bytes)) as zin:
        existing = {i.filename: zin.read(i.filename) for i in zin.infolist()}

    rels_xml = existing.get(rels_path)
    if rels_xml is None:
        rels_xml = (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>'
        )
    rel_frags = "".join(
        f'<Relationship Id="{rid}" Type="{rtype}" Target="{target}"/>'
        for rid, rtype, target in rels
    ).encode("utf-8")
    rels_xml = rels_xml.replace(b"</Relationships>", rel_frags + b"</Relationships>")
    existing[rels_path] = rels_xml

    ct_xml = existing[ct_path]
    ct_frags = "".join(
        f'<Override PartName="{pn}" ContentType="{ct}"/>' for pn, ct in content_types
    ).encode("utf-8")
    ct_xml = ct_xml.replace(b"</Types>", ct_frags + b"</Types>")
    existing[ct_path] = ct_xml

    existing.update(parts)

    with zipfile.ZipFile(output, "w") as zout:
        for name, data in existing.items():
            zout.writestr(name, data)
    return output.getvalue()


# ============================ E05(1) 编号体系 ============================

_NUMBERING_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    f'<w:numbering xmlns:w="{W_NS}">'
    '<w:abstractNum w:abstractNumId="0">'
    '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>'
    '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2"/></w:lvl>'
    '</w:abstractNum>'
    '<w:abstractNum w:abstractNumId="1">'
    '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="\u2022"/></w:lvl>'
    '</w:abstractNum>'
    '<w:abstractNum w:abstractNumId="2">'
    '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>'
    '</w:abstractNum>'
    '<w:num w:numId="10"><w:abstractNumId w:val="0"/></w:num>'
    '<w:num w:numId="11"><w:abstractNumId w:val="1"/></w:num>'
    '<w:num w:numId="12"><w:abstractNumId w:val="2"/>'
    '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="5"/></w:lvlOverride></w:num>'
    '</w:numbering>'
).encode("utf-8")


def _docx_with_numbering(body_paragraphs: bytes) -> bytes:
    doc = Document()
    doc.add_paragraph("ANCHORPARA")
    base = _save(doc)
    # 锚段形如 <w:p><w:r><w:t>ANCHORPARA</w:t></w:r></w:p>。把整个 run 替换成
    # "</w:p>{各完整段落}<w:p>"：锚段被拆成一头一尾两个空 <w:p></w:p>（无害），
    # 中间注入的每个段落都自成完整 <w:p>...</w:p>，整体 XML 保持平衡。
    base = _rewrite(
        base, "word/document.xml",
        b"<w:r><w:t>ANCHORPARA</w:t></w:r>",
        b"</w:p>" + body_paragraphs + b"<w:p>",
    )
    return _add_parts(
        base,
        parts={"word/numbering.xml": _NUMBERING_XML},
        rels=[("rIdNum99", f"{R_NS}/numbering", "numbering.xml")],
        content_types=[("/word/numbering.xml",
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml")],
    )


def _num_para(num_id: str, ilvl: str, text: str) -> bytes:
    return (
        f'<w:p><w:pPr><w:numPr><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/></w:numPr></w:pPr>'
        f'<w:r><w:t>{text}</w:t></w:r></w:p>'
    ).encode("utf-8")


def test_e05_multilevel_chapter_numbering_symbol_and_category():
    body = _num_para("10", "0", "First") + _num_para("10", "1", "First-Sub")
    file_bytes = _docx_with_numbering(body)
    blocks = _blocks(file_bytes)
    numbered = [b for b in blocks if b.metadata and b.metadata.get("numbering")]
    assert len(numbered) == 2

    top = numbered[0].metadata["numbering"]
    assert top["symbol"] == "1."
    assert top["category"] == "chapter"
    assert top["numbering_fidelity"] == "exact"

    sub = numbered[1].metadata["numbering"]
    assert sub["symbol"] == "1.1"
    assert sub["ilvl"] == 1
    assert sub["category"] == "chapter"


def test_e05_bullet_list_is_body_list_not_chapter():
    body = _num_para("11", "0", "bullet item")
    blocks = _blocks(_docx_with_numbering(body))
    numbered = [b for b in blocks if b.metadata and b.metadata.get("numbering")]
    assert numbered
    num = numbered[0].metadata["numbering"]
    assert num["num_fmt"] == "bullet"
    assert num["symbol"] == "\u2022"
    assert num["category"] == "list"


def test_e05_start_override_restarts_numbering_at_five():
    body = _num_para("12", "0", "restarted")
    blocks = _blocks(_docx_with_numbering(body))
    numbered = [b for b in blocks if b.metadata and b.metadata.get("numbering")]
    assert numbered
    num = numbered[0].metadata["numbering"]
    assert num["symbol"] == "5."
    assert num["restarted"] is True


def test_e05_automatic_numbering_reset_across_two_top_level_groups():
    body = (
        _num_para("10", "0", "T1")
        + _num_para("10", "1", "T1-a")
        + _num_para("10", "0", "T2")
        + _num_para("10", "1", "T2-a")
    )
    blocks = _blocks(_docx_with_numbering(body))
    syms = [b.metadata["numbering"]["symbol"] for b in blocks
            if b.metadata and b.metadata.get("numbering")]
    assert syms == ["1.", "1.1", "2.", "2.1"]


# ============================ E05(2) 脚注 / 尾注 / 书签 / 内部引用 ============================

_FOOTNOTES_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    f'<w:footnotes xmlns:w="{W_NS}">'
    '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
    '<w:footnote w:id="1"><w:p><w:r><w:t>FOOTNOTE BODY ONE</w:t></w:r></w:p></w:footnote>'
    '</w:footnotes>'
).encode("utf-8")


def test_e05_footnote_reference_position_and_body_are_recoverable():
    doc = Document()
    doc.add_paragraph("CITEHERE")
    base = _save(doc)
    base = _rewrite(
        base, "word/document.xml",
        b"<w:r><w:t>CITEHERE</w:t></w:r>",
        b'<w:r><w:t>CITEHERE</w:t></w:r>'
        b'<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
        b'<w:footnoteReference w:id="1"/></w:r>',
    )
    file_bytes = _add_parts(
        base,
        parts={"word/footnotes.xml": _FOOTNOTES_XML},
        rels=[("rIdFn99", f"{R_NS}/footnotes", "footnotes.xml")],
        content_types=[("/word/footnotes.xml",
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml")],
    )
    block = next(b for b in _blocks(file_bytes) if b.type == "text" and "CITEHERE" in (b.text or ""))
    refs = block.metadata.get("note_refs")
    assert refs, "脚注引用位置必须保留在引用所在段"
    ref = refs[0]
    assert ref["kind"] == "footnote"
    assert ref["note_id"] == "1"
    assert ref["marker"] == 1
    assert ref["text"] == "FOOTNOTE BODY ONE"


def test_e05_bookmark_and_internal_reference_positions_survive():
    doc = Document()
    doc.add_paragraph("TARGETPARA")
    doc.add_paragraph("REFPARA")
    base = _save(doc)
    base = _rewrite(
        base, "word/document.xml",
        b"<w:r><w:t>TARGETPARA</w:t></w:r>",
        b'<w:bookmarkStart w:id="1" w:name="sec1"/>'
        b'<w:r><w:t>TARGETPARA</w:t></w:r><w:bookmarkEnd w:id="1"/>',
    )
    base = _rewrite(
        base, "word/document.xml",
        b"<w:r><w:t>REFPARA</w:t></w:r>",
        b'<w:hyperlink w:anchor="sec1"><w:r><w:t>see section</w:t></w:r></w:hyperlink>',
    )
    blocks = _blocks(base)
    target = next(b for b in blocks if "TARGETPARA" in (b.text or ""))
    assert target.metadata.get("bookmarks")
    assert any(bm["name"] == "sec1" for bm in target.metadata["bookmarks"])

    ref = next(b for b in blocks if b.metadata and b.metadata.get("internal_refs"))
    assert any(r["anchor"] == "sec1" for r in ref.metadata["internal_refs"])


# ============================ E05(3) cell 内嵌内容 + 嵌套表入块 ============================

def test_e05_nested_table_becomes_its_own_block_with_span_relationship():
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=2)
    inner.cell(0, 0).text = "INNERA"
    inner.cell(0, 1).text = "INNERB"

    blocks = _blocks(_save(doc))
    tables = [b for b in blocks if b.type == "table"]
    assert len(tables) == 2, "内层表必须真正入块（不再被丢弃）"

    outer_block, inner_block = tables[0], tables[1]
    outer_cell = outer_block.metadata["cells"][0]
    assert outer_cell["nested_table_block_ids"] == [inner_block.id]
    assert inner_block.metadata["parent_table"]["block_id"] == outer_block.id
    assert inner_block.metadata["parent_table"]["row"] == 0
    assert inner_block.metadata["parent_table"]["col"] == 0
    assert inner_block.metadata["table_depth"] == 1
    assert "INNERA" in inner_block.text and "INNERB" in inner_block.text
    assert "INNERA" not in (outer_block.text or "")


def test_e05_cell_image_is_recorded_as_inline_content():
    import base64

    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
    )
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    para = table.cell(0, 0).paragraphs[0]
    para.add_run("cap ")
    para.add_run().add_picture(BytesIO(png))

    table_block = next(b for b in _blocks(_save(doc)) if b.type == "table")
    cell = table_block.metadata["cells"][0]
    assert cell.get("content"), "cell 内的图片应记入 content"
    assert any(item["kind"] == "image" for item in cell["content"])
    assert table_block.metadata["has_inline_content"] is True


def test_e05_cell_formula_content_carries_fidelity():
    doc = Document()
    doc.add_table(rows=1, cols=1)
    base = _save(doc)
    frac = (
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        '<m:f><m:num><m:r><m:t>a</m:t></m:r></m:num>'
        '<m:den><m:r><m:t>b</m:t></m:r></m:den></m:f></m:oMath>'
    ).encode("utf-8")
    base = _rewrite(
        base, "word/document.xml",
        b"</w:tc>",
        b"<w:p>" + frac + b"</w:p></w:tc>",
    )
    table_block = next(b for b in _blocks(base) if b.type == "table")
    cell = table_block.metadata["cells"][0]
    formula_items = [i for i in (cell.get("content") or []) if i["kind"] == "formula"]
    assert formula_items
    assert formula_items[0]["latex"] == "\\frac{a}{b}"
    assert formula_items[0]["formula_fidelity"] == "exact"


# ============================ E05(4) OMML formula_fidelity ============================

def _docx_with_body_omml(math_xml: str) -> bytes:
    doc = Document()
    doc.add_paragraph("MATHANCHOR")
    base = _save(doc)
    # 把锚段的 run 替换为 "</w:p>{含公式的段落}<w:p>"，保持 XML 平衡。
    return _rewrite(
        base, "word/document.xml",
        b"<w:r><w:t>MATHANCHOR</w:t></w:r>",
        b'</w:p><w:p>' + math_xml.encode("utf-8") + b'</w:p><w:p>',
    )


def test_e05_simple_formula_is_marked_exact():
    math = (
        '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        '<m:oMath><m:sSup><m:e><m:r><m:t>x</m:t></m:r></m:e>'
        '<m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:oMath></m:oMathPara>'
    )
    block = next(b for b in _blocks(_docx_with_body_omml(math)) if b.type == "formula")
    assert block.text == "x^{2}"
    assert block.metadata["formula_fidelity"] == "exact"


def test_e05_matrix_formula_is_marked_lossy_and_keeps_source():
    math = (
        '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        '<m:oMath><m:m><m:mr><m:e><m:r><m:t>1</m:t></m:r></m:e>'
        '<m:e><m:r><m:t>2</m:t></m:r></m:e></m:mr></m:m></m:oMath></m:oMathPara>'
    )
    block = next(b for b in _blocks(_docx_with_body_omml(math)) if b.type == "formula")
    assert block.metadata["formula_fidelity"] == "lossy"
    assert block.metadata.get("source_omml")


def test_e05_formula_and_text_in_same_paragraph():
    doc = Document()
    doc.add_paragraph("prefix  suffix")
    base = _save(doc)
    frac = (
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        '<m:f><m:num><m:r><m:t>a</m:t></m:r></m:num>'
        '<m:den><m:r><m:t>b</m:t></m:r></m:den></m:f></m:oMath>'
    ).encode("utf-8")
    base = _rewrite(
        base, "word/document.xml",
        b"<w:t>prefix  suffix</w:t>",
        b'<w:t xml:space="preserve">prefix </w:t></w:r>' + frac + b'<w:r><w:t xml:space="preserve"> suffix</w:t>',
    )
    block = next(b for b in _blocks(base) if b.type == "text")
    assert block.text == "prefix $\\frac{a}{b}$ suffix"


# ============================ E05(5) SmartArt / OLE / 未知 XML ============================

def test_e05_unknown_smartart_becomes_visible_placeholder_not_html_comment():
    doc = Document()
    doc.add_paragraph("KEEPME")
    base = _save(doc)
    smartart = (
        b'</w:p>'
        b'<dgm:relIds xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" '
        b'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        b'r:dm="rId1"/>'
        b'<w:p>'
    )
    base = _rewrite(
        base, "word/document.xml",
        b"<w:r><w:t>KEEPME</w:t></w:r>",
        b"<w:r><w:t>KEEPME</w:t></w:r>" + smartart,
    )
    blocks = _blocks(base)
    placeholder = next(
        (b for b in blocks if b.metadata and b.metadata.get("source") == "unknown_element"),
        None,
    )
    assert placeholder is not None, "未知元素必须以可见占位块保留"
    assert placeholder.text and placeholder.text.startswith("\u3010")
    assert "<!--" not in placeholder.text
    assert placeholder.metadata["unknown_kind"] == "smartart"

    prov = DocxProvider()
    prov.extract_from_bytes(base)
    codes = [s.code for s in prov.drain_skips()]
    assert "docx_unknown_element" in codes


# ============================ E05(6) P02 code 合并逐字符等价 ============================

def _reference_merge(code_texts: list) -> str:
    """参考实现：连续 code 段的合并应等价于 '\\n'.join(各段文本)。"""
    return "\n".join(code_texts)


def test_p02_consecutive_code_blocks_merge_is_char_equivalent_to_join():
    doc = Document()
    code_lines = ["line one alpha", "line two beta", "line three gamma"]
    for line in code_lines:
        para = doc.add_paragraph()
        run = para.add_run(line)
        run.font.name = "Consolas"

    blocks = _blocks(_save(doc))
    code_blocks = [b for b in blocks if b.type == "code"]
    assert len(code_blocks) == 1, "连续 code 段应合并成一块"
    merged = code_blocks[0]
    assert merged.text == _reference_merge(code_lines)

    atoms = merged.metadata["atoms"]
    for atom, expected in zip(atoms, code_lines):
        assert merged.text[atom["char_start"]:atom["char_end"]] == expected


def test_p02_single_code_block_is_left_intact():
    doc = Document()
    para = doc.add_paragraph()
    run = para.add_run("only('one')")
    run.font.name = "Consolas"

    code_blocks = [b for b in _blocks(_save(doc)) if b.type == "code"]
    assert len(code_blocks) == 1
    assert code_blocks[0].text == "only('one')"
